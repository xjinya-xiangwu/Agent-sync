import sys
from pptx import Presentation
from pptx.util import Emu

path = sys.argv[1]
prs = Presentation(path)
W, H = prs.slide_width, prs.slide_height
print(f"canvas: {W/914400:.2f} x {H/914400:.2f} in, slides: {len(prs.slides.__iter__.__self__._sldIdLst)}")

issues = 0
for si, slide in enumerate(prs.slides, 1):
    boxes = []
    for sh in slide.shapes:
        if sh.left is None:
            continue
        l, t, w, h = sh.left, sh.top, sh.width, sh.height
        if l < -10000 or t < -10000 or l + w > W + 10000 or t + h > H + 10000:
            print(f"OUT-OF-BOUNDS: {sh.shape_type} at ({l/914400:.2f},{t/914400:.2f}) {w/914400:.2f}x{h/914400:.2f}")
            issues += 1
        if sh.has_text_frame and sh.text_frame.text.strip():
            boxes.append((l, t, w, h, sh.text_frame.text.strip()[:14]))
    for i in range(len(boxes)):
        for j in range(i + 1, len(boxes)):
            a, b = boxes[i], boxes[j]
            ox = min(a[0]+a[2], b[0]+b[2]) - max(a[0], b[0])
            oy = min(a[1]+a[3], b[1]+b[3]) - max(a[1], b[1])
            if ox > 91440 and oy > 45720:  # >0.1in x >0.05in overlap
                print(f"OVERLAP: '{a[4]}' vs '{b[4]}' ({ox/914400:.2f}x{oy/914400:.2f} in)")
                issues += 1
    # 占位符残留
    for sh in slide.shapes:
        if sh.has_text_frame:
            for bad in ("xxx", "lorem", "TODO", "[insert", "[必填]"):
                if bad in sh.text_frame.text:
                    print(f"PLACEHOLDER LEFT: {bad}")
                    issues += 1

print("QA issues:", issues)
